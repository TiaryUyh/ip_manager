import csv
import io
import ipaddress
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from .models import IpRecord

EXPORT_HEADERS = [
    "IP地址",
    "网段",
    "状态",
    "主机名",
    "MAC地址",
    "使用人",
    "部门",
    "位置",
    "备注",
]

FIELD_ALIASES = {
    "ip地址": "ip_address",
    "ip": "ip_address",
    "ip_address": "ip_address",
    "ip地址(必填)": "ip_address",
    "网段": "cidr",
    "cidr": "cidr",
    "状态": "status",
    "status": "status",
    "主机名": "hostname",
    "hostname": "hostname",
    "mac地址": "mac_address",
    "mac": "mac_address",
    "mac_address": "mac_address",
    "使用人": "assignee",
    "assignee": "assignee",
    "部门": "department",
    "department": "department",
    "位置": "location",
    "location": "location",
    "备注": "description",
    "description": "description",
}

STATUS_ALIASES = {
    "available": "available",
    "空闲": "available",
    "空闲可用": "available",
    "in_use": "in_use",
    "使用中": "in_use",
    "占用": "in_use",
    "reserved": "reserved",
    "预留": "reserved",
    "保留": "reserved",
    "disabled": "disabled",
    "停用": "disabled",
    "禁用": "disabled",
}

STATUS_LABELS = {
    "available": "空闲",
    "in_use": "使用中",
    "reserved": "预留",
    "disabled": "停用",
}

VALID_STATUSES = set(STATUS_LABELS.keys())


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def record_to_row(record: IpRecord) -> list[str]:
    return [
        record.ip_address or "",
        record.cidr or "",
        STATUS_LABELS.get(record.status, record.status or ""),
        record.hostname or "",
        record.mac_address or "",
        record.assignee or "",
        record.department or "",
        record.location or "",
        record.description or "",
    ]


def build_csv_bytes(records: list[IpRecord]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    for record in records:
        writer.writerow(record_to_row(record))
    # Excel 友好：加 BOM
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def build_xlsx_bytes(records: list[IpRecord]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IP记录"
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for record in records:
        sheet.append(record_to_row(record))
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            max_length = max(max_length, len(str(cell.value or "")))
        sheet.column_dimensions[column_letter].width = min(max_length + 4, 40)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_template_xlsx_bytes() -> bytes:
    sample = [
        [
            "192.168.1.10",
            "192.168.1.0/24",
            "使用中",
            "pc-dev",
            "AA:BB:CC:DD:EE:01",
            "张三",
            "研发部",
            "3F-A01",
            "示例数据，导入前可删除",
        ]
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IP记录"
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sample:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_template_csv_bytes() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    writer.writerow(
        [
            "192.168.1.10",
            "192.168.1.0/24",
            "使用中",
            "pc-dev",
            "AA:BB:CC:DD:EE:01",
            "张三",
            "研发部",
            "3F-A01",
            "示例数据，导入前可删除",
        ]
    )
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def _normalize_header(header: Any) -> str | None:
    key = _clean(header)
    if not key:
        return None
    return FIELD_ALIASES.get(key.lower().replace(" ", ""))


def _parse_status(value: Any) -> str:
    text = _clean(value) or "in_use"
    mapped = STATUS_ALIASES.get(text.lower()) or STATUS_ALIASES.get(text)
    if not mapped:
        raise ValueError(f"状态无效：{text}（可用：空闲/使用中/预留/停用）")
    return mapped


def _validate_ip(value: str) -> str:
    try:
        ip = ipaddress.ip_address(value)
    except ValueError as exc:
        raise ValueError("请输入有效的 IPv4 地址") from exc
    if ip.version != 4:
        raise ValueError("仅支持 IPv4 地址")
    return str(ip)


def _rows_from_csv(content: bytes) -> list[dict[str, str | None]]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    headers = [_normalize_header(h) for h in rows[0]]
    if "ip_address" not in headers:
        raise ValueError("缺少必填列：IP地址")
    result = []
    for row in rows[1:]:
        if not any(_clean(cell) for cell in row):
            continue
        item: dict[str, str | None] = {}
        for index, field in enumerate(headers):
            if not field:
                continue
            item[field] = _clean(row[index]) if index < len(row) else None
        result.append(item)
    return result


def _rows_from_xlsx(content: bytes) -> list[dict[str, str | None]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_normalize_header(h) for h in rows[0]]
    if "ip_address" not in headers:
        raise ValueError("缺少必填列：IP地址")
    result = []
    for row in rows[1:]:
        if not row or not any(_clean(cell) for cell in row):
            continue
        item: dict[str, str | None] = {}
        for index, field in enumerate(headers):
            if not field:
                continue
            item[field] = _clean(row[index]) if index < len(row) else None
        result.append(item)
    return result


def parse_import_file(filename: str, content: bytes) -> list[dict[str, str | None]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _rows_from_csv(content)
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _rows_from_xlsx(content)
    raise ValueError("仅支持 .csv 或 .xlsx 文件")


def import_rows(db: Session, rows: list[dict[str, str | None]]) -> dict[str, Any]:
    created = 0
    updated = 0
    failed: list[dict[str, Any]] = []

    for index, row in enumerate(rows, start=2):
        try:
            ip_raw = row.get("ip_address")
            if not ip_raw:
                raise ValueError("IP地址不能为空")
            ip_address = _validate_ip(ip_raw)
            status = _parse_status(row.get("status"))
            payload = {
                "ip_address": ip_address,
                "cidr": row.get("cidr"),
                "status": status,
                "hostname": row.get("hostname"),
                "mac_address": row.get("mac_address"),
                "assignee": row.get("assignee"),
                "department": row.get("department"),
                "location": row.get("location"),
                "description": row.get("description"),
            }

            existing = (
                db.query(IpRecord).filter(IpRecord.ip_address == ip_address).first()
            )
            if existing:
                for key, value in payload.items():
                    if key == "ip_address":
                        continue
                    setattr(existing, key, value)
                updated += 1
            else:
                db.add(IpRecord(**payload))
                created += 1
        except Exception as exc:  # noqa: BLE001 - 汇总行级错误
            failed.append({"row": index, "ip": row.get("ip_address"), "error": str(exc)})

    if created or updated:
        db.commit()

    return {
        "total": len(rows),
        "created": created,
        "updated": updated,
        "failed": len(failed),
        "errors": failed[:50],
    }
